"""Check that the committed sample documents match what the generator produces.

Builds a fresh copy into a temporary folder, then compares content, not
bytes: every PDF page's text and the PDF metadata (where one of the hidden
injections lives), every spreadsheet cell, and every answer-key file. If a
committed document has been edited by hand, or its hidden test content has
changed, this exits with an error and lists the differences.

Run from the repository root:

    python3 scripts/verify_sample_data.py

Requires: the packages in requirements-dev.txt.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
COMMITTED = ROOT / "sample-data"


def pdf_content(path: Path) -> dict:
    reader = PdfReader(path)
    meta = reader.metadata or {}
    return {
        "metadata": {key: str(meta.get(key)) for key in ("/Title", "/Author", "/Subject")},
        "pages": [page.extract_text() for page in reader.pages],
    }


def xlsx_content(path: Path) -> dict:
    wb = load_workbook(path)
    props = wb.properties
    return {
        "properties": {"creator": props.creator, "created": str(props.created)},
        "sheets": {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets},
    }


def content(path: Path):
    if path.suffix == ".pdf":
        return pdf_content(path)
    if path.suffix == ".xlsx":
        return xlsx_content(path)
    return path.read_text()


def generated_files(folder: Path) -> set[Path]:
    return {p.relative_to(folder) for p in folder.rglob("*") if p.is_file() and p.suffix in {".pdf", ".xlsx", ".json"}}


def main() -> int:
    with tempfile.TemporaryDirectory() as tmp:
        fresh = Path(tmp)
        subprocess.run([sys.executable, str(ROOT / "scripts" / "generate_sample_data.py"), "--out", str(fresh)],
                       check=True, stdout=subprocess.DEVNULL)
        expected, committed = generated_files(fresh), generated_files(COMMITTED)
        problems = [f"missing from sample-data/: {p}" for p in sorted(expected - committed)]
        problems += [f"not produced by the generator: {p}" for p in sorted(committed - expected)]
        for rel in sorted(expected & committed):
            if content(fresh / rel) != content(COMMITTED / rel):
                problems.append(f"content differs: {rel}")
    if problems:
        print("Sample data does not match the generator:")
        for problem in problems:
            print(f"  - {problem}")
        print("Run: python3 scripts/generate_sample_data.py, review the changes, and commit them.")
        return 1
    print(f"Sample data matches the generator ({len(expected)} files checked by content).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
