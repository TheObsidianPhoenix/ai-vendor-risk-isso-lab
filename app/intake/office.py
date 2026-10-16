"""XLSX and DOCX checks. Both formats are zip archives of XML files.

Order matters: the archive itself is checked (size, entry count, compression
ratio, macros, embedded objects, XML declarations) before any library parses
its contents, so a hostile archive is rejected without being opened by
openpyxl or the XML parser.
"""

from __future__ import annotations

import io
import re
import zipfile
import xml.etree.ElementTree as ET

from .limits import Limits
from .patterns import instruction_like
from .result import Flag, Rejected, Segment

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CORE_FIELDS = ("title", "subject", "creator", "keywords", "description", "lastModifiedBy")
WHITE = {"FFFFFF", "FFFFFFFF"}


# ---------------------------------------------------------------------------
# Archive safety (shared)
# ---------------------------------------------------------------------------

def open_archive(data: bytes, limits: Limits) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise Rejected("Damaged or invalid Office file") from None
    entries = archive.infolist()
    if len(entries) > limits.max_zip_entries:
        raise Rejected(f"Archive has {len(entries)} parts; the limit is {limits.max_zip_entries}")
    total = 0
    for entry in entries:
        name = entry.filename.lower()
        if entry.flag_bits & 0x1:
            raise Rejected("Encrypted content inside the file")
        if name.endswith("vbaproject.bin") or "macro" in name:
            raise Rejected("File contains macros")
        if "/embeddings/" in name or "activex" in name or name.endswith("oleobject.bin"):
            raise Rejected("File contains embedded objects or ActiveX controls")
        total += entry.file_size
        if entry.compress_size and entry.file_size / entry.compress_size > limits.max_compression_ratio:
            raise Rejected(f"Part {entry.filename} expands {entry.file_size // entry.compress_size}x: possible zip bomb")
    if total > limits.max_uncompressed_bytes:
        raise Rejected(f"File expands to {total // (1024 * 1024)} MB; the limit is "
                       f"{limits.max_uncompressed_bytes // (1024 * 1024)} MB")
    for entry in entries:
        if entry.filename.lower().endswith((".xml", ".rels")) and b"<!DOCTYPE" in archive.read(entry).upper():
            # Office files never need a DTD; one is how XML-bomb and external-entity attacks start
            raise Rejected(f"Part {entry.filename} declares a DTD, which Office files never need")
    return archive


def metadata_flags(archive: zipfile.ZipFile, limits: Limits) -> list[Flag]:
    try:
        root = ET.fromstring(archive.read("docProps/core.xml"))
    except KeyError:
        return []
    flags = []
    for element in root:
        field = element.tag.rsplit("}", 1)[-1]
        if field in CORE_FIELDS and element.text and instruction_like(element.text):
            flags.append(Flag("metadata_instruction", f"metadata {field}",
                              "Document metadata reads like an instruction to an AI",
                              element.text[: limits.max_excerpt_chars]))
    return flags


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def scan_xlsx(data: bytes, limits: Limits) -> tuple[list[Flag], list[Segment]]:
    archive = open_archive(data, limits)
    names = [n.lower() for n in archive.namelist()]
    if "xl/workbook.xml" not in names:
        raise Rejected("Not a valid Excel workbook")
    flags = metadata_flags(archive, limits)
    if any(n.startswith("xl/externallinks/") for n in names):
        flags.append(Flag("external_link", "workbook", "Workbook links to data in another file"))

    from openpyxl import load_workbook  # Imported here: only reached after the archive passed every check
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    segments: list[Segment] = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            flags.append(Flag("hidden_sheet", f"sheet '{sheet.title}'", f"Sheet is {sheet.sheet_state}"))
        hidden_rows = {r for r, dim in sheet.row_dimensions.items() if dim.hidden}
        hidden_cols = {c for c, dim in sheet.column_dimensions.items() if dim.hidden}
        if hidden_rows or hidden_cols:
            flags.append(Flag("hidden_rows_or_columns", f"sheet '{sheet.title}'",
                              f"{len(hidden_rows)} hidden rows, {len(hidden_cols)} hidden columns"))
        for row in sheet.iter_rows():
            visible_values = []
            for cell in row:
                if cell.value is None or str(cell.value).strip() == "":
                    continue
                text = str(cell.value)
                location = f"sheet '{sheet.title}' cell {cell.coordinate}"
                color = cell.font.color.rgb if cell.font and cell.font.color and cell.font.color.type == "rgb" else None
                hidden = (sheet.sheet_state != "visible" or cell.row in hidden_rows
                          or cell.column_letter in hidden_cols or (isinstance(color, str) and color.upper() in WHITE))
                if hidden:
                    if isinstance(color, str) and color.upper() in WHITE:
                        flags.append(Flag("hidden_text", location, "White text, invisible on a white cell",
                                          text[: limits.max_excerpt_chars]))
                    continue  # Hidden content is never passed to the model
                phrase = instruction_like(text)
                if phrase:
                    flags.append(Flag("instruction_like_text", location, "Cell text reads like an instruction to an AI", phrase))
                visible_values.append(text)
            if visible_values:
                segments.append(Segment(f"sheet '{sheet.title}' row {row[0].row}", " | ".join(visible_values)))
    if sum(len(s.text) for s in segments) < limits.min_extracted_chars:
        raise Rejected("Workbook has no readable content")
    return flags, segments


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def scan_docx(data: bytes, limits: Limits) -> tuple[list[Flag], list[Segment]]:
    archive = open_archive(data, limits)
    try:
        body = ET.fromstring(archive.read("word/document.xml"))
    except KeyError:
        raise Rejected("Not a valid Word document") from None
    flags = metadata_flags(archive, limits)
    if any(n.lower().startswith("word/_rels/") and b"TargetMode=\"External\"" in archive.read(n)
           for n in archive.namelist()):
        flags.append(Flag("external_link", "document", "Document references an external resource"))

    segments: list[Segment] = []
    for number, paragraph in enumerate(body.iter(f"{W}p"), start=1):
        location = f"paragraph {number}"
        visible, hidden = [], []
        for run in paragraph.iter(f"{W}r"):
            text = "".join(t.text or "" for t in run.iter(f"{W}t"))
            if not text:
                continue
            reason = _hidden_run_reason(run)
            (hidden if reason else visible).append((reason, text))
        for reason, text in hidden:
            flags.append(Flag("hidden_text", location, reason, text[: limits.max_excerpt_chars]))
        line = "".join(text for _, text in visible).strip()
        if line:
            phrase = instruction_like(line)
            if phrase:
                flags.append(Flag("instruction_like_text", location, "Visible text reads like an instruction to an AI", phrase))
            segments.append(Segment(location, line))
    if sum(len(s.text) for s in segments) < limits.min_extracted_chars:
        raise Rejected("Document has no readable text")
    return flags, segments


def _hidden_run_reason(run) -> str | None:
    props = run.find(f"{W}rPr")
    if props is None:
        return None
    if props.find(f"{W}vanish") is not None or props.find(f"{W}webHidden") is not None:
        return "Text marked hidden in Word"
    color = props.find(f"{W}color")
    if color is not None and (color.get(f"{W}val") or "").upper() in WHITE:
        return "White text, invisible on a white page"
    size = props.find(f"{W}sz")
    if size is not None and re.fullmatch(r"\d+", size.get(f"{W}val") or "") and int(size.get(f"{W}val")) < 8:
        return f"Text at {int(size.get(f'{W}val')) / 2:.1f} pt, too small to read"
    return None
